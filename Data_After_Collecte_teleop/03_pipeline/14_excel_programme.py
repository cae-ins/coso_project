"""14 — Produire le classeur Excel « analyse programme » (variante Arthur).

Entrée : 04_resultats/analyse_programme/indicateurs_programme.json
Sortie : 04_resultats/analyse_programme/Analyse_Programme_COSO.xlsx

Classeur SÉPARÉ du classeur officiel des 21 indicateurs. Une feuille par bloc
(P1 à P10), plus Methode et MAF_cellules. Fermer le classeur avant de relancer.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ICI = Path(__file__).resolve().parent
RACINE = ICI.parent
ENTREE = RACINE / "04_resultats" / "analyse_programme" / "indicateurs_programme.json"
SORTIE = RACINE / "04_resultats" / "analyse_programme" / "Analyse_Programme_COSO.xlsx"

BLEU = "17324D"
FOND = PatternFill("solid", fgColor=BLEU)
GRIS = PatternFill("solid", fgColor="EEF2F6")
TITRE = Font(bold=True, color=BLEU, size=12)
SOUS = Font(bold=True, size=10)
NOTE = Font(italic=True, size=9, color="555555")
BLANC = Font(color="FFFFFF", bold=True)

FEUILLES = {
    "P1": "P1_Filieres",
    "P2": "P2_Saturation",
    "P3": "P3_Formation",
    "P4": "P4_Vulnerabilite",
    "P5": "P5_Metiers_Gap",
    "P6": "P6_Ecart_MAF",
    "P7": "P7_Geographie",
    "P8": "P8_Revenus",
    "P9": "P9_Croisements",
    "P10": "P10_Entonnoir",
}


def _ecrire_tableau(ws, ligne, tab):
    ws.cell(row=ligne, column=1, value=f"{tab['cle']} — {tab['titre']}").font = TITRE
    ligne += 1
    if tab.get("note"):
        ws.cell(row=ligne, column=1, value=tab["note"]).font = NOTE
        ws.cell(row=ligne, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne,
                       end_column=max(2, len(tab["entetes"])))
        ligne += 2
    entetes = tab["entetes"]
    for j, e in enumerate(entetes, start=1):
        c = ws.cell(row=ligne, column=j, value=e)
        c.font = BLANC
        c.fill = FOND
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ligne += 1
    for i, rang in enumerate(tab["lignes"]):
        for j, v in enumerate(rang, start=1):
            c = ws.cell(row=ligne, column=j, value=v)
            if i == len(tab["lignes"]) - 1 or str(rang[0]).lower() == "total":
                c.font = SOUS
            if str(rang[0]).lower() == "total":
                c.fill = GRIS
        ligne += 1
    return ligne + 1


def _largeurs(ws, tab):
    for j, e in enumerate(tab["entetes"], start=1):
        col = get_column_letter(j)
        largeur = max(12, min(46, len(str(e)) + 4))
        if j == 1:
            largeur = max(largeur, 34)
        ws.column_dimensions[col].width = largeur


def produire(entree: Path = ENTREE, sortie: Path = SORTIE) -> dict:
    payload = json.loads(entree.read_text(encoding="utf-8"))
    meta = payload["metadata"]
    wb = Workbook()

    # — Feuille Methode —
    ws = wb.active
    ws.title = "Methode"
    ws.cell(row=1, column=1, value="Analyse programme — volet exploratoire (variante Arthur)").font = Font(bold=True, size=14, color=BLEU)
    lignes = [
        ("", ""),
        ("Date de calcul", meta["date_calcul"]),
        ("Population", f"{meta['population']} candidats (notre population, étape 07)"),
        ("Base", meta["base"]),
        ("Mode sector_match", meta["mode_sector_match"]),
        ("Source MAF", meta["source_maf"]),
        ("", ""),
        ("Avertissement", meta["avertissement"]),
        ("", ""),
        ("Blocs", ""),
    ]
    for cle, lib in meta["blocs"].items():
        lignes.append((cle, lib))
    for i, (a, b) in enumerate(lignes, start=2):
        ws.cell(row=i, column=1, value=a).font = SOUS
        ws.cell(row=i, column=2, value=b).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 100
    ws.cell(row=len(lignes) + 4, column=1,
            value="Ce classeur est séparé du rapport officiel des 21 indicateurs. "
                  "Il contient des scores composites et une donnée externe (SONET-CI) "
                  "et ne doit pas être présenté comme un résultat de suivi-évaluation validé.").font = NOTE

    # — Feuilles par bloc —
    for bloc, nom in FEUILLES.items():
        tables = [t for t in payload["tableaux"] if t["bloc"] == bloc]
        if not tables:
            continue
        wsb = wb.create_sheet(nom)
        ligne = 1
        wsb.cell(row=1, column=1, value=meta["blocs"].get(bloc, bloc)).font = Font(bold=True, size=13, color=BLEU)
        ligne = 3
        for tab in tables:
            ligne = _ecrire_tableau(wsb, ligne, tab)
            _largeurs(wsb, tab)

    # — Feuille MAF brut —
    maf = payload.get("maf", {})
    wsm = wb.create_sheet("MAF_cellules")
    wsm.cell(row=1, column=1, value="Matrice MAF — cellules métier × région").font = Font(bold=True, size=13, color=BLEU)
    wsm.cell(row=2, column=1, value=maf.get("source", {}).get("reference", "")).font = NOTE
    entetes = ["Code", "Métier", "Région", "Demande", "Offre (MAF)", "Ratio", "Feu"]
    for j, e in enumerate(entetes, start=1):
        c = wsm.cell(row=4, column=j, value=e)
        c.font = BLANC
        c.fill = FOND
    for i, r in enumerate(maf.get("cellules", []), start=5):
        wsm.cell(row=i, column=1, value=r["trade_code"])
        wsm.cell(row=i, column=2, value=r["metier"])
        wsm.cell(row=i, column=3, value=r["region"])
        wsm.cell(row=i, column=4, value=r["demande"])
        wsm.cell(row=i, column=5, value=r["supply"])
        wsm.cell(row=i, column=6, value=r["ratio"])
        wsm.cell(row=i, column=7, value=r["flag"])
    for col, w in zip("ABCDEFG", [7, 32, 16, 11, 12, 10, 9]):
        wsm.column_dimensions[col].width = w

    sortie.parent.mkdir(parents=True, exist_ok=True)
    wb.save(sortie)
    return {"sortie": str(sortie), "feuilles": wb.sheetnames,
            "tableaux": len(payload["tableaux"])}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--entree", type=Path, default=ENTREE)
    parser.add_argument("--sortie", type=Path, default=SORTIE)
    args = parser.parse_args()
    print(json.dumps(produire(args.entree, args.sortie), ensure_ascii=False, indent=2))
