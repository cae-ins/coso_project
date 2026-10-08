"""11 — Contrôles de cohérence des livrables d'indicateurs.

Contrôle, sans modifier aucun fichier :
  · la base apurée et l'audit de l'étape 07 (population, signalements, PII) ;
  · indicateurs.json (structure, additivité des sous-groupes, repères de valeurs) ;
  · le classeur Excel (feuilles, tailles) ;
  · le rapport Word (sections, intégrité du paquet, pas de données nominatives).

Sortie console : une ligne [OK] ou [ERREUR] par contrôle, code retour 1 si échec.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path

import pandas as pd
from docx import Document
from openpyxl import load_workbook

ICI = Path(__file__).resolve().parent
RACINE = ICI.parent
BASE = RACINE / "04_resultats" / "traitement" / "base_apuree.csv"
META = RACINE / "04_resultats" / "traitement" / "meta_apurement.json"
JSON_INDI = RACINE / "04_resultats" / "indicateurs" / "indicateurs.json"
XLSX = RACINE / "04_resultats" / "indicateurs" / "Indicateurs_COSO.xlsx"
DOCX = RACINE / "04_resultats" / "indicateurs" / "Rapport_indicateurs_COSO.docx"

PII_COLONNES = re.compile(r"^(nom|telephone|B7|R1A|R2A|R3B|R4|R6)", re.I)
PII_TEXTE = re.compile(r"\b0[1-9]\d{8}\b|[\w.+-]+@[\w-]+\.[\w.]+")
FEUILLES = ["Methode", "Synthese_globale", "Modalites", "Sous_groupes",
            "Couverture_21", "Non_calculables", "Meta"]
SECTIONS = ["1. Objet et sources", "2. Chaîne de traitement", "3. Règles de calcul",
            "4. Couverture des 21 indicateurs", "5. Résultats par bloc",
            "6. Codages des questions ouvertes", "7. Sous-groupes", "8. Limites",
            "9. Reproductibilité"]

echecs: list[str] = []
controles = 0


def controler(nom: str, condition: bool, detail: str = "") -> bool:
    global controles
    controles += 1
    if condition:
        print(f"[OK] {nom}")
    else:
        message = f"{nom}" + (f" — {detail}" if detail else "")
        print(f"[ERREUR] {message}")
        echecs.append(message)
    return bool(condition)


def controler_base() -> dict:
    meta = json.loads(META.read_text(encoding="utf-8"))
    d = pd.read_csv(BASE, low_memory=False)

    controler("Étape 07 : 962 entretiens en source", meta["n_source"] == 962,
              f"n_source={meta['n_source']}")
    controler("Étape 07 : 664 après complétude",
              meta["n_apres_completude"] == 664, str(meta["n_apres_completude"]))
    controler("Étape 07 : 663 candidats analysés", meta["n_apures"] == 663,
              str(meta["n_apures"]))
    controler("Étape 07 : 71 signalements", meta["n_signalements"] == 71,
              str(meta["n_signalements"]))
    controler("Étape 07 : 61 entretiens signalés",
              meta["n_entretiens_signales"] == 61, str(meta["n_entretiens_signales"]))
    controler("Base apurée : 663 lignes, identifiants uniques",
              len(d) == 663 and d.interview__key.is_unique,
              f"{len(d)} lignes")
    controler("Base apurée : population homogeneous",
              d.classe_completude.isin(["1-COMPLET", "2-QUASI_COMPLET"]).all()
              and d.consentement.eq(1).all())
    controler("Base apurée : aucune colonne nominative",
              not any(PII_COLONNES.match(str(c)) for c in d.columns),
              str([c for c in d.columns if PII_COLONNES.match(str(c))]))
    controler("Base apurée : 7 âges hors plage 15–120 ans (laissés vides)",
              int((~d.age_valide.between(15, 120)).sum()) == 7)
    duree = d.duree_entretien_min[d.duree_valide]
    controler("Durée : 515 entretiens dans [1 ; 600] min, médiane 23,4",
              len(duree) == 515 and abs(duree.median() - 23.4) <= 0.1,
              f"n={len(duree)}, médiane={duree.median()}")
    controler("Emploi : 600 en emploi, 63 sans emploi",
              int(d.emploi_calcule.eq(1).sum()) == 600
              and int(d.emploi_calcule.eq(0).sum()) == 63)
    sans = d[d.emploi_calcule.eq(0)]
    controler("Emploi : F1 sur ses 7 modalités officielles",
              d.F1.dropna().isin([1, 2, 3, 4, 5, 7, 9]).all(),
              str(sorted(d.F1.dropna().unique())))
    controler("Emploi : J0 renseigné et ∈ {1,2} pour les 63 sans emploi",
              len(sans) == 63 and sans.J0.isin([1, 2]).all())
    return meta


def _lignes(payload: dict) -> list[dict]:
    return payload["resultats"]


def controler_structure(payload: dict) -> None:
    lignes = _lignes(payload)
    controler("indicateurs.json : lignes avec groupe et bloc renseignés",
              all(l.get("niveau") and l.get("bloc", "").startswith("D") for l in lignes))
    cles = [(l["niveau"], l["district"], l["region"], l["sexe"], l["age_groupe"],
             l["milieu"], l["code"], l["mesure"], l["modalite"]) for l in lignes]
    controler("indicateurs.json : clés uniques", len(cles) == len(set(cles)),
              f"{len(cles) - len(set(cles))} doublons")
    hors_agregation = [l for l in lignes if l["effectif"] is not None
                       and l["denominateur"] is not None
                       and (l["effectif"] > l["denominateur"]
                            or (l["denominateur"] > l["n_groupe"]
                                and l["code"] != "D2.2"))]
    controler("indicateurs.json : effectif ≤ dénominateur ≤ taille du groupe "
              "(D2.2 mesure par personnes du ménage)", not hors_agregation,
              str([(l["code"], l["niveau"]) for l in hors_agregation[:5]]))
    controler("indicateurs.json : sans_reponse ≥ 0",
              all(l["sans_reponse"] is None or l["sans_reponse"] >= 0 for l in lignes))
    controler("indicateurs.json : aucune donnée nominative dans les libellés",
              not any(PII_TEXTE.search(str(l.get(c) or ""))
                      for l in lignes
                      for c in ("modalite", "mesure", "precision", "variables")))
    couverture = payload["couverture"]
    statuts = pd.Series([e["statut"] for e in couverture]).value_counts().to_dict()
    controler("Couverture : 21 indicateurs, 3 en attente, 8 indisponibles",
              len(couverture) == 21 and len({e["code"] for e in couverture}) == 21
              and statuts.get("en attente") == 3
              and statuts.get("indisponible") == 8, str(statuts))
    controler("Couverture et non_calculables cohérents",
              {e["code"] for e in couverture if e["statut"] in
               ("en attente", "indisponible")}
              == {e["code"] for e in payload["non_calculables"]},
              f"{len(payload['non_calculables'])} non calculables")
    controler("indicateurs.json : 23 groupes de désagrégation",
              len({(l["niveau"], l["district"], l["region"], l["sexe"],
                    l["age_groupe"], l["milieu"]) for l in lignes}) == 23)


def controler_additivite(payload: dict) -> None:
    lignes = [l for l in _lignes(payload) if l["denominateur"] is not None]
    df = pd.DataFrame(lignes)
    df["groupe"] = (df.niveau + "|" + df.district + "|" + df.region + "|"
                    + df.sexe + "|" + df.age_groupe + "|" + df.milieu)
    groupes = df[df.niveau != "Ensemble"].drop_duplicates(
        ["niveau", "groupe"] if "groupe" in df.columns else ["niveau"])
    tailles = groupes.groupby("niveau")["n_groupe"].sum()
    controler("Sous-groupes : chaque niveau totalise 663",
              all(tailles.get(n, 0) == 663 for n in
                  ["District", "Région", "Sexe", "Âge", "Milieu"]),
              str(tailles.to_dict()))
    for niveau in ["District", "Région", "Sexe", "Âge", "Milieu"]:
        sub = df[df.niveau.isin([niveau, "Ensemble"])]
        pivot = sub.pivot_table(
            index=["code", "mesure", "modalite"], columns="niveau",
            values=["effectif", "denominateur"], aggfunc="sum", fill_value=0)
        manquants = [c for c in ("effectif", "denominateur") if c not in pivot.columns
                     or niveau not in pivot[c]]
        ok = not manquants
        ecart = 0
        if ok:
            ecart = int((pivot[("effectif", niveau)] != pivot[("effectif", "Ensemble")]).sum())
            ecart += int((pivot[("denominateur", niveau)]
                          != pivot[("denominateur", "Ensemble")]).sum())
        controler(f"Additivité {niveau} : effectifs et dénominateurs somment "
                  f"à l'ensemble", ok and ecart == 0,
                  f"lignes divergentes : {ecart}")


def controler_reperes(payload: dict) -> None:
    e = { (l["code"], l["modalite"]): l for l in _lignes(payload)
          if l["niveau"] == "Ensemble"}
    attendus = [
        (("D1.1", "Femme"), 227, 663, "%"),
        (("D3.1", "En emploi"), 600, 663, "%"),
        (("D5.4", ""), 23, 63, "%"),
        (("D8.1", ""), 135, 663, "%"),
        (("D10.1", ""), 271, 663, "%"),
        (("D10.9", ""), 378, 612, "%"),
        (("D9.5", ""), 111, 663, "%"),
        (("D9.6", ""), 520, 663, "%"),
        (("D9.7", ""), 32, 663, "%"),
        (("D6.2", ""), 296, 326, "%"),
    ]
    for cle, effet, denom, unite in attendus:
        l = e.get(cle)
        controler(f"Repère {cle[0]} {cle[1] or '(mesure unique)'} : "
                  f"{effet}/{denom} {unite}",
                  l is not None and l["effectif"] == effet
                  and l["denominateur"] == denom,
                  f"obtenu {l and l['effectif']}/{l and l['denominateur']}")
    medianes = [(("D4.1", ""), 50000), (("D4.4", ""), 38500),
                (("D3.8", ""), 45), (("D1.2", ""), 31),
                (("D11.3", ""), 23.4), (("D2.16", ""), 4),
                (("D7.4", ""), 25000), (("D8.9", ""), 150000)]
    for cle, valeur in medianes:
        l = e.get(cle)
        controler(f"Médiane {cle[0]} : {valeur}",
                  l is not None and abs(l["valeur"] - valeur) < max(0.05,
                                                                    abs(valeur) * 1e-6),
                  f"obtenu {l and l['valeur']}")


def controler_excel(payload: dict) -> None:
    if not controler("Excel : fichier présent", XLSX.exists()):
        return
    wb = load_workbook(XLSX, read_only=False)
    controler("Excel : 7 feuilles attendues", wb.sheetnames == FEUILLES,
              str(wb.sheetnames))
    n_lignes = len(payload["resultats"])
    controler("Excel : Modalites = 9017 lignes (en-tête comprise)",
              wb["Modalites"].max_row == n_lignes + 1,
              str(wb["Modalites"].max_row))
    controler("Excel : Sous_groupes = 24 lignes (en-tête comprise)",
              wb["Sous_groupes"].max_row == 24, str(wb["Sous_groupes"].max_row))
    controler("Excel : Couverture_21 = 22 lignes (en-tête comprise)",
              wb["Couverture_21"].max_row == 22, str(wb["Couverture_21"].max_row))
    controler("Excel : tableaux structurants présents",
              all(wb[f].tables for f in ("Synthese_globale", "Modalites",
                                         "Sous_groupes", "Couverture_21",
                                         "Non_calculables")))


def controler_word() -> None:
    if not controler("Word : fichier présent", DOCX.exists()):
        return
    with zipfile.ZipFile(DOCX) as paquet:
        controler("Word : paquet DOCX intègre", paquet.testzip() is None)
        controler("Word : parties XML principales présentes",
                  {"[Content_Types].xml", "word/document.xml", "word/styles.xml"}
                  <= set(paquet.namelist()))
    doc = Document(DOCX)
    textes = [p.text.strip() for p in doc.paragraphs]
    for section in SECTIONS:
        controler(f"Word : section « {section} »", section in textes)
    controler("Word : titre présent",
              any(p.style.name == "Title" and "indicateurs COSO" in p.text
                  for p in doc.paragraphs))
    controler("Word : au moins 15 tableaux", len(doc.tables) >= 15,
              str(len(doc.tables)))
    corps = "\n".join(textes)
    corps += "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    controler("Word : population annoncée = 663", "663" in corps)
    controler("Word : aucune donnée nominative", not PII_TEXTE.search(corps))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    args = parser.parse_args()

    if not controler("Base apurée présente", BASE.exists()):
        print(f"\n{controles} contrôles, {len(echecs)} échec(s).")
        sys.exit(1)
    controler_base()
    charge = json.loads(JSON_INDI.read_text(encoding="utf-8"))
    controler_structure(charge)
    controler_additivite(charge)
    controler_reperes(charge)
    controler_excel(charge)
    controler_word()

    print(f"\n{controles} contrôles, {len(echecs)} échec(s).")
    if echecs:
        for message in echecs:
            print(f"  ✗ {message}")
        sys.exit(1)
