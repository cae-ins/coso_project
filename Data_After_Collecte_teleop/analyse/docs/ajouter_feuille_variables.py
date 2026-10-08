#!/usr/bin/env python3
"""
Ajoute la feuille "Variables" au classeur QUESTIONNAIRE_VF_SANS_DOUBLON_VF.xlsx.

La feuille "Variables" est le dictionnaire lisible de la feuille "Donnees" :
une ligne par colonne, dans l'ordre exact du tableau, avec la section du
questionnaire, l'intitule de la question (recupere dans le HTML), le type,
la completude, un exemple de valeur, les modalites et une description.

Entrees :
  - analyse/docs/Original Questionnaire_COSO_V5.html   (intitules + types + modalites)
  - analyse/base/QUESTIONNAIRE_VF_SANS_DOUBLON_VF.xlsx (codebook + valeurs + donnees)

Sortie : le meme .xlsx, avec la feuille "Variables" ajoutee en deuxieme position.

Usage : python3 ajouter_feuille_variables.py
"""
import re
from html import unescape

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HTML_SRC = ("/Users/macbookair/Desktop/CAE/COSO/COSO-repo/coso_project/"
            "Data_After_Collecte_teleop/analyse/docs/Original Questionnaire_COSO_V5.html")
XLSX = ("/Users/macbookair/Desktop/CAE/COSO/COSO-repo/coso_project/"
        "Data_After_Collecte_teleop/analyse/base/QUESTIONNAIRE_VF_SANS_DOUBLON_VF.xlsx")

FEUILLE_VARIABLES = "Variables"
SECTION_SYSTEME = "SYSTÈME / CALCULÉ"

TYPE_FR = {
    "text": "Texte",
    "numeric: integer": "Numérique (entier)",
    "numeric: decimal": "Numérique (décimal)",
    "single-select": "Choix unique",
    "single-select: cascading": "Choix unique (en cascade)",
    "multi-select": "Choix multiple",
    "date": "Date",
    "date: current time": "Date et heure",
    "double": "Numérique (décimal)",
    "long": "Numérique (entier)",
    "string": "Texte",
    "bool": "Booléen (0/1)",
}

NOTES = {
    "interview__key": "Clé unique de l'entretien (district-n°), générée par Survey Solutions.",
    "interview__id": "Identifiant technique de l'entretien (UUID).",
    "interview__status": "Statut de l'entretien dans Survey Solutions (voir modalités).",
    "assignment__id": "Identifiant de la mission (assignment) attribuée à l'agent.",
    "sssys_irnd": "Aléatoire système (sssys_irnd) utilisé pour l'assignation des entretiens.",
    "has__errors": "1 si l'entretien contient des erreurs de validation, 0 sinon.",
    "age": "Âge calculé = 2026 - année de naissance (C2A_annee).",
    "n_attendu": "Nombre de questions censées être posées (logique de saut du questionnaire).",
    "n_rempli": "Nombre de questions effectivement renseignées.",
    "taux_completude": "Proportion 0-1 des questions attendues renseignées (n_rempli / n_attendu).",
    "classe_completude": "Classe de complétude : 1-COMPLET · 2-QUASI_COMPLET · 3-MOYEN · 4-INCOMPLET · 5-VIDE.",
    "doublon_de": "Motif d'exclusion si la fiche a été écartée comme doublon (vide sinon).",
    "doublon_classe": "classe_completude de l'entretien jumeau écarté (vide sinon).",
    "a_retraiter": "1 si la fiche reste à arbitrer (complétude 3-MOYEN, 4-INCOMPLET ou 5-VIDE et non déjà écartée).",
    "consentement": "Consentement du répondant : 1 = Oui, 0 = Non, vide = non renseigné (voir consentement_lbl).",
    "consentement_lbl": "Libellé du consentement : OUI / NON / NA (NA = non renseigné).",
    "EMPLOYE": "Variable calculée : 1 si le répondant exerce une activité rémunérée (E1==1 ou E2==1 ou (E3==1 et E3A==1)), 0 sinon.",
}

LARGEURS = [5, 18, 34, 62, 20, 11, 15, 13, 14, 26, 58, 55]


def texte_brut(fragment):
    fragment = re.sub(r"<br\s*/?>", " ", fragment)
    fragment = re.sub(r"<[^>]+>", " ", fragment)
    return re.sub(r"\s+", " ", unescape(fragment)).strip()


def blocs_div(html, classe):
    """Renvoie les blocs <div class="classe"> ... </div> a div equilibres."""
    blocs = []
    for m in re.finditer(r'<div class="%s"' % re.escape(classe), html):
        debut, profondeur, curseur = m.start(), 0, m.start()
        while True:
            trouve = re.search(r"<div\b|</div>", html[curseur:])
            if not trouve:
                break
            jeton = trouve.group(0)
            curseur += trouve.start()
            profondeur += -1 if jeton == "</div>" else 1
            curseur += 4 if jeton == "</div>" else len(jeton)
            if profondeur == 0:
                blocs.append(html[debut:curseur])
                break
    return blocs


def lire_questionnaire(chemin):
    """Extrait {variable: section, titre, type, options} depuis le HTML."""
    html = open(chemin, encoding="utf-8", errors="replace").read()
    html = html[html.find("</style>"):]

    entetes = [
        (m.start(), texte_brut(m.group(1)))
        for m in re.finditer(
            r'<div class="section_header">.*?<h2[^>]*>(.*?)</h2>', html, re.S)
    ]

    questions = {}
    for bloc in blocs_div(html, "question-container"):
        position = html.find(bloc)
        titre = re.search(r'class="question-title"[^>]*>(.*?)</div>', bloc, re.S)
        qtype = re.search(
            r'class="question-meta">.*?class="type">\s*(.*?)\s*</div>', bloc, re.S)
        variable = re.search(r'class="variable_name">(.*?)</div>', bloc, re.S)
        expression = re.search(
            r'class="variable-expression">(.*?)</div>', bloc, re.S)
        if not variable:
            continue

        section = ""
        for debut, nom in entetes:
            if debut < position:
                section = nom
            else:
                break

        options = {}
        for valeur, libelle in zip(
            re.findall(r'<div class="option-value">\s*(.*?)\s*</div>', bloc, re.S),
            re.findall(r'<label for="[^"]*">(.*?)</label>', bloc, re.S),
        ):
            cle = texte_brut(valeur)
            if cle:
                options[cle] = texte_brut(libelle)

        brut_type = texte_brut(qtype.group(1)) if qtype else ""
        intitule = ""
        if titre and 'class="type"' not in titre.group(1):
            intitule = texte_brut(titre.group(1))
        questions[texte_brut(variable.group(1))] = {
            "section": section,
            "titre": intitule,
            "type": TYPE_FR.get(brut_type, brut_type),
            "options": options,
            "expression": texte_brut(expression.group(1)) if expression else "",
        }

    for variable, info in list(questions.items()):
        if info["type"] != "Choix multiple":
            continue
        for code, libelle in info["options"].items():
            numero = int(code) if code.isdigit() else code
            questions[f"{variable}__{numero}"] = {
                "section": info["section"],
                "titre": f'{info["titre"]} — modalité {code} : {libelle}',
                "type": "Choix multiple (0/1)",
                "options": {"1": "cochée", "0": "non cochée"},
                "expression": "",
            }

    return questions


def principale():
    questions = lire_questionnaire(HTML_SRC)
    print(f"Questions extraites du questionnaire HTML : {len(questions)}")

    classeur = load_workbook(XLSX)

    codebook = {}
    for ligne in classeur["Codebook"].iter_rows(min_row=2, values_only=True):
        if ligne[0]:
            codebook[ligne[0]] = {
                "type": ligne[3],
                "n_ok": ligne[4],
                "n_manquants": ligne[5],
                "pct": ligne[6],
            }

    modalites = {}
    for variable, code, libelle in classeur["Valeurs"].iter_rows(min_row=2, values_only=True):
        if variable:
            modalites.setdefault(variable, []).append(f"{code} = {libelle}")

    donnees = classeur["Donnees"]
    colonnes = [c.value for c in donnees[1]]

    exemples = {}
    for ligne in donnees.iter_rows(min_row=2, values_only=True):
        for position, valeur in enumerate(ligne):
            colonne = colonnes[position]
            if colonne not in exemples and valeur not in (None, ""):
                exemples[colonne] = str(valeur)[:60]

    if FEUILLE_VARIABLES in classeur.sheetnames:
        del classeur[FEUILLE_VARIABLES]
    feuille = classeur.create_sheet(FEUILLE_VARIABLES, 1)

    entetes = ["N°", "Variable", "Section", "Question", "Type", "Format",
               "N_non_manquants", "N_manquants", "Pct_manquants",
               "Exemple", "Modalités", "Description"]
    remplis = PatternFill("solid", start_color="1F4E78")
    police_entete = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    police_corps = Font(name="Arial", size=10)
    trait = Side(style="thin", color="B0B0B0")
    bordure = Border(left=trait, right=trait, top=trait, bottom=trait)
    retour = Alignment(wrap_text=True, vertical="top")

    for colonne, entete in enumerate(entetes, start=1):
        cellule = feuille.cell(row=1, column=colonne, value=entete)
        cellule.fill = remplis
        cellule.font = police_entete
        cellule.border = bordure
        cellule.alignment = Alignment(vertical="center")

    sans_intitule = []
    for numero, colonne in enumerate(colonnes, start=1):
        info = questions.get(colonne)
        if info:
            section, question, qtype = info["section"], info["titre"], info["type"]
            options_html = info["options"]
        else:
            section, question, qtype, options_html = SECTION_SYSTEME, "", "", {}

        description = NOTES.get(colonne, "")
        expression = info["expression"] if info else ""
        if not question:
            if expression:
                question = "Variable calculée"
            elif description:
                question = description.split(" (")[0].rstrip(".")
        if not description and expression:
            description = "Variable calculée : " + expression[:300]
        if not info:
            sans_intitule.append(colonne)

        liste_modalites = modalites.get(colonne) or [
            f"{code} = {libelle}" for code, libelle in options_html.items()
        ]
        texte_modalites = " ; ".join(liste_modalites)
        if len(texte_modalites) > 900:
            texte_modalites = texte_modalites[:900] + " …"

        stats = codebook.get(colonne, {})
        valeurs = [
            numero,
            colonne,
            section,
            question,
            qtype,
            stats.get("type"),
            stats.get("n_ok"),
            stats.get("n_manquants"),
            stats.get("pct"),
            exemples.get(colonne),
            texte_modalites or None,
            description or None,
        ]

        for position, valeur in enumerate(valeurs, start=1):
            cellule = feuille.cell(row=numero + 1, column=position, value=valeur)
            cellule.font = police_corps
            cellule.border = bordure
            cellule.alignment = retour

    for position, largeur in enumerate(LARGEURS, start=1):
        feuille.column_dimensions[get_column_letter(position)].width = largeur
    feuille.freeze_panes = "C2"
    feuille.auto_filter.ref = feuille.dimensions

    classeur.save(XLSX)

    print(f"Feuille '{FEUILLE_VARIABLES}' ecrite : {len(colonnes)} variables")
    print(f"  dont {len(sans_intitule)} sans intitule dans le HTML "
          f"(decrites dans la colonne Description)")
    print("OK ->", XLSX)


if __name__ == "__main__":
    principale()
