#!/usr/bin/env python3
"""
Crée bases/QUESTIONNAIRE_VF_SANS_DOUBLON_VF.xlsx contenant :
  - feuille "Codebook" : toutes les variables avec leur définition (libellé),
    type, complétude et valeurs possibles ;
  - feuille "Valeurs"  : dictionnaire des codes/valeurs (format long) ;
  - feuille "Donnees"  : les 962 lignes de données brutes.
"""
import os
import re

import numpy as np
import pandas as pd
import pyreadstat

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC = "/Users/macbookair/Desktop/CAE/COSO/COSO-repo/coso_project/Data_After_Collecte_teleop/resultats/QUESTIONNAIRE_VF_SANS_DOUBLON_VF.dta"
OUT = "/Users/macbookair/Desktop/CAE/COSO/COSO-repo/coso_project/Data_After_Collecte_teleop/bases/QUESTIONNAIRE_VF_SANS_DOUBLON_VF.xlsx"

# ---------------------------------------------------------------------------
# Lecture du .dta
# ---------------------------------------------------------------------------
print("Lecture du .dta ...")
df, meta = pyreadstat.read_dta(SRC, apply_value_formats=False)
labels_var = meta.column_names_to_labels
labels_vals = meta.variable_value_labels

def module_de(col):
    m = re.match(r"^([A-Z]{1,3})", col)
    return m.group(1) if m else "AUTRE"

# ---------------------------------------------------------------------------
# Feuille 1 : Codebook (définitions des variables)
# ---------------------------------------------------------------------------
rows = []
for col in df.columns:
    vals = labels_vals.get(col, {})
    valeurs = "; ".join(f"{k} - {v}" for k, v in vals.items()) or ""
    rows.append({
        "Variable": col,
        "Module": module_de(col),
        "Definition": labels_var.get(col, ""),
        "Type": str(df[col].dtype),
        "N_non_manquants": int(df[col].notna().sum()),
        "N_manquants": int(df[col].isna().sum()),
        "Pct_manquants": round(100 * df[col].isna().mean(), 1),
        "Valeurs_possibles": valeurs,
    })
codebook = pd.DataFrame(rows)

# ---------------------------------------------------------------------------
# Feuille 2 : dictionnaire des valeurs (long)
# ---------------------------------------------------------------------------
lab_rows = []
for c, d in labels_vals.items():
    for k, v in d.items():
        lab_rows.append({"Variable": c, "Code": k, "Libelle": v})
valeurs = pd.DataFrame(lab_rows, columns=["Variable", "Code", "Libelle"])

# ---------------------------------------------------------------------------
# Écriture Excel
# ---------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", start_color="1F4E78")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="Arial", size=10)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

def give_sheet(ws, frame):
    # entêtes
    for j, col in enumerate(frame.columns, start=1):
        cell = ws.cell(row=1, column=j, value=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        cell.border = BORDER
    # corps
    for i, (_, r) in enumerate(frame.iterrows(), start=2):
        for j, col in enumerate(frame.columns, start=1):
            v = r[col]
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = WRAP
    # largeurs
    for j, col in enumerate(frame.columns, start=1):
        w = min(max(12, *frame[col].astype(str).str.len().dropna()), 60) if len(frame) else 12
        ws.column_dimensions[get_column_letter(j)].width = w + 2
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

ws_cb = wb.active
ws_cb.title = "Codebook"
give_sheet(ws_cb, codebook)

ws_va = wb.create_sheet("Valeurs")
give_sheet(ws_va, valeurs)

ws_dd = wb.create_sheet("Donnees")
for j, col in enumerate(df.columns, start=1):
    cell = ws_dd.cell(row=1, column=j, value=col)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
ws_dd.append(df.values.tolist()[0] if len(df) else [])
for i, row in enumerate(df.itertuples(index=False, name=None), start=2):
    for j, v in enumerate(row, start=1):
        cell = ws_dd.cell(row=i, column=j, value=None if pd.isna(v) else v)
        cell.font = BODY_FONT
        cell.border = BORDER
ws_dd.freeze_panes = "B2"
ws_dd.auto_filter.ref = ws_dd.dimensions

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("OK ->", OUT)
print(f"Feuilles : {wb.sheetnames} | Codebook : {len(codebook)} variables | "
      f"Valeurs : {len(valeurs)} codes | Donnees : {df.shape}")



